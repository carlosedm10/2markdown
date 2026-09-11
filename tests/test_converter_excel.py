"""Tests for Excel conversion (twomarkdown.converter.excel)."""

from pathlib import Path
from unittest.mock import patch

import pytest

from twomarkdown.converter.excel import (
    ExcelConversionError,
    convert_xlsx,
    rows_to_markdown_table,
)


class TestRowsToMarkdownTable:
    def test_rows_to_markdown_table_basic(self) -> None:
        table = rows_to_markdown_table([["A", "B"], ["1", "2"]])
        assert "| A | B |" in table
        assert "| --- | --- |" in table
        assert "| 1 | 2 |" in table


class TestConvertXlsx:
    def test_convert_xlsx_reads_sheets(self, tmp_path: Path) -> None:
        pytest.importorskip("openpyxl")
        from openpyxl import Workbook

        xlsx_path = tmp_path / "book.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws["A1"] = "Name"
        ws["B1"] = "Count"
        ws["A2"] = "alpha"
        ws["B2"] = 3
        wb.create_sheet("Empty")
        wb.save(xlsx_path)

        markdown = convert_xlsx(xlsx_path)

        assert "## Sheet: Data" in markdown
        assert "| Name | Count |" in markdown
        assert "| alpha | 3 |" in markdown
        assert "Empty" not in markdown

    def test_convert_xlsx_wide_sheet_uses_csv_block(self, tmp_path: Path) -> None:
        pytest.importorskip("openpyxl")
        from openpyxl import Workbook

        xlsx_path = tmp_path / "wide.xlsx"
        wb = Workbook()
        ws = wb.active
        for col in range(1, 32):
            ws.cell(row=1, column=col, value=f"c{col}")
        wb.save(xlsx_path)

        markdown = convert_xlsx(xlsx_path)

        assert "```csv" in markdown
        assert "c31" in markdown
        assert "| c1 |" not in markdown

    def test_convert_xlsx_missing_openpyxl_raises(self, tmp_path: Path) -> None:
        xlsx_path = tmp_path / "book.xlsx"
        xlsx_path.write_bytes(b"not-a-real-xlsx")

        with patch.dict("sys.modules", {"openpyxl": None}):
            with pytest.raises(ExcelConversionError, match="openpyxl"):
                convert_xlsx(xlsx_path)

    def test_convert_xlsx_preserves_formulas_alongside_values(
        self, tmp_path: Path
    ) -> None:
        """convert_xlsx() — formula cells keep the formula in a companion block."""
        pytest.importorskip("openpyxl")
        from openpyxl import Workbook

        xlsx_path = tmp_path / "formulas.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Calc"
        ws["A1"] = "Min"
        ws["B1"] = "Max"
        ws["C1"] = "Range"
        ws["A2"] = 30
        ws["B2"] = 35
        ws["C2"] = "=B2-A2"
        wb.save(xlsx_path)

        markdown = convert_xlsx(xlsx_path)

        assert "| 30 | 35 |" in markdown
        assert "**Fórmulas:**" in markdown
        assert "`C2 = B2-A2`" in markdown
        assert "=B2-A2" not in markdown.split("**Fórmulas:**")[0]

    def test_convert_xlsx_no_formula_block_when_sheet_has_no_formulas(
        self, tmp_path: Path
    ) -> None:
        """convert_xlsx() — a sheet with no formulas emits no Fórmulas block."""
        pytest.importorskip("openpyxl")
        from openpyxl import Workbook

        xlsx_path = tmp_path / "plain.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Plain"
        ws["A1"] = "Name"
        ws["A2"] = "alpha"
        wb.save(xlsx_path)

        markdown = convert_xlsx(xlsx_path)

        assert "Fórmulas" not in markdown

    def test_convert_xlsx_caps_formula_block_and_reports_omitted_count(
        self, tmp_path: Path
    ) -> None:
        """convert_xlsx() — caps the formula block and reports the omitted count."""
        pytest.importorskip("openpyxl")
        from openpyxl import Workbook

        from twomarkdown.converter.excel import MAX_FORMULAS_PER_SHEET

        xlsx_path = tmp_path / "many_formulas.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Many"
        total = MAX_FORMULAS_PER_SHEET + 10
        for row in range(1, total + 1):
            ws.cell(row=row, column=1, value=row)
            ws.cell(row=row, column=2, value=f"=A{row}*2")
        wb.save(xlsx_path)

        markdown = convert_xlsx(xlsx_path)

        assert markdown.count("`B") == MAX_FORMULAS_PER_SHEET
        assert "+10 fórmulas más omitidas" in markdown

    def test_convert_xlsx_propagates_merged_header_value(
        self, tmp_path: Path
    ) -> None:
        """convert_xlsx() — a merged range repeats its top-left value, not blanks."""
        pytest.importorskip("openpyxl")
        from openpyxl import Workbook

        xlsx_path = tmp_path / "merged.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Merged"
        ws["A1"] = "Quarterly results"
        ws.merge_cells("A1:C1")
        ws["A2"] = "Q1"
        ws["B2"] = "Q2"
        ws["C2"] = "Q3"
        wb.save(xlsx_path)

        markdown = convert_xlsx(xlsx_path)

        header = "| Quarterly results | Quarterly results | Quarterly results |"
        assert header in markdown

    def test_convert_xlsx_empty_workbook_raises(self, tmp_path: Path) -> None:
        """convert_xlsx() — a workbook with no cell data raises."""
        pytest.importorskip("openpyxl")
        from openpyxl import Workbook

        xlsx_path = tmp_path / "empty.xlsx"
        wb = Workbook()
        wb.save(xlsx_path)

        with pytest.raises(ExcelConversionError, match="no sheet data"):
            convert_xlsx(xlsx_path)
