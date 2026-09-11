"""Convert Excel .xlsx workbooks to markdown."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import TYPE_CHECKING

if TYPE_CHECKING:
    from openpyxl.worksheet.worksheet import Worksheet

MAX_GFM_COLUMNS = 30

# A sheet with thousands of formulas would otherwise dump an unreadable wall
# of text after the table; cap the companion block and say what was cut.
MAX_FORMULAS_PER_SHEET = 50


class ExcelConversionError(Exception):
    """Raised when Excel conversion cannot be performed."""


def rows_to_markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    lines: list[str] = []
    for i, row in enumerate(normalized):
        line = "| " + " | ".join(cell.replace("|", "\\|") for cell in row) + " |"
        lines.append(line)
        if i == 0:
            lines.append("| " + " | ".join("---" for _ in row) + " |")
    return "\n".join(lines)


def _rows_to_csv_block(rows: list[list[str]]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    for row in rows:
        writer.writerow(row)
    return f"```csv\n{buffer.getvalue().rstrip()}\n```"


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").strip()


def _sheet_has_data(rows: list[list[str]]) -> bool:
    return any(any(cell for cell in row) for row in rows)


def _values_by_coordinate(sheet: Worksheet) -> dict[tuple[int, int], object]:
    """Map (row, column) -> cell value, then flatten merged ranges onto it.

    Merged cells only carry a value on their top-left member; every other
    cell in the range is a MergedCell with value None. Propagating the
    top-left value across the range keeps merged header rows from
    extracting as blanks.
    """
    values: dict[tuple[int, int], object] = {}
    for row in sheet.iter_rows():
        for cell in row:
            values[(cell.row, cell.column)] = cell.value

    for merged_range in sheet.merged_cells.ranges:
        top_left = values.get((merged_range.min_row, merged_range.min_col))
        for r in range(merged_range.min_row, merged_range.max_row + 1):
            for c in range(merged_range.min_col, merged_range.max_col + 1):
                values[(r, c)] = top_left

    return values


def _sheet_rows(sheet: Worksheet) -> list[list[str]]:
    values = _values_by_coordinate(sheet)
    rows: list[list[str]] = []
    for r in range(sheet.min_row, sheet.max_row + 1):
        col_range = range(sheet.min_column, sheet.max_column + 1)
        row_values = [values.get((r, c)) for c in col_range]
        str_row = [_cell_to_str(v) for v in row_values]
        if any(str_row):
            rows.append(str_row)
    return rows


def _sheet_formulas(sheet: Worksheet) -> list[tuple[str, str]]:
    """Collect (coordinate, formula) pairs, ordered top-to-bottom then left-to-right."""
    formulas: list[tuple[str, str]] = []
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type != "f" or not isinstance(cell.value, str):
                continue
            text = cell.value
            if text.startswith("="):
                text = text[1:]
            formulas.append((cell.coordinate, text))
    return formulas


def _formula_block(formulas: list[tuple[str, str]]) -> str:
    if not formulas:
        return ""
    shown = formulas[:MAX_FORMULAS_PER_SHEET]
    omitted = len(formulas) - len(shown)
    entries = " · ".join(f"`{coord} = {text}`" for coord, text in shown)
    line = f"**Fórmulas:** {entries}"
    if omitted > 0:
        line += f" · _(+{omitted} fórmulas más omitidas)_"
    return line


def convert_xlsx(path: Path) -> str:
    """Convert an .xlsx workbook to markdown with one section per sheet."""
    try:
        import openpyxl
    except ImportError as exc:
        raise ExcelConversionError(
            "openpyxl is required for Excel conversion: pip install openpyxl"
        ) from exc

    path = path.resolve()
    sections: list[str] = []

    # Two full parses, not read_only: read_only worksheets don't expose
    # merged_cells at all, and a single data_only=True load discards every
    # formula (it only returns the last cached value). Correctness over
    # speed here, per the task.
    try:
        values_workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    except Exception as exc:
        raise ExcelConversionError(f"xlsx read failed: {exc}") from exc

    try:
        formulas_workbook = openpyxl.load_workbook(
            path, read_only=False, data_only=False
        )
    except Exception as exc:
        values_workbook.close()
        raise ExcelConversionError(f"xlsx read failed: {exc}") from exc

    try:
        for sheet in values_workbook.worksheets:
            rows = _sheet_rows(sheet)

            if not _sheet_has_data(rows):
                continue

            sections.append(f"## Sheet: {sheet.title}")

            width = max(len(row) for row in rows) if rows else 0
            if width > MAX_GFM_COLUMNS:
                sections.append(_rows_to_csv_block(rows))
            else:
                sections.append(rows_to_markdown_table(rows))
            sections.append("")

            if sheet.title in formulas_workbook.sheetnames:
                formulas = _sheet_formulas(formulas_workbook[sheet.title])
                block = _formula_block(formulas)
                if block:
                    sections.append(block)
                    sections.append("")
    finally:
        values_workbook.close()
        formulas_workbook.close()

    text = "\n".join(sections).strip()
    if not text:
        raise ExcelConversionError("xlsx: no sheet data extracted")
    return text
